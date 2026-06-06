# Cài đặt thư viện trước khi chạy: pip install python-pptx
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title_text, content_text):
    slide_layout = prs.slide_layouts[1] # Sử dụng layout 'Title and Content'
    slide = prs.slides.add_slide(slide_layout)
    
    # Thiết lập tiêu đề
    title_shape = slide.shapes.title
    title_shape.text = title_text
    
    # Thiết lập nội dung
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.word_wrap = True
    tf.text = content_text

# Slide 1: Trang tiêu đề
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.shapes.placeholders[1]
title.text = "Dự đoán chính xác tương tác thuốc và tác dụng phụ thuốc bằng mô hình lai CNN-BILSTM"
subtitle.text = "Đại học Bách khoa Hà Nội - Trường CNTT & TT\nNhóm 19 - Lớp IT3190\nThành viên: Nguyễn Đức Quý, Nguyễn Văn Phú Thái, Nguyễn Ngọc Lê Minh, Đào Xuân Hùng, Lê Đình Doanh"

# Slide 2: Lý do chọn đề tài
add_slide("1. Lý do chọn đề tài", 
          "• Tương tác thuốc (DDI) có thể gây tác dụng phụ nguy hiểm hoặc tử vong.\n"
          "• Thử nghiệm lâm sàng truyền thống tốn kém thời gian và chi phí.\n"
          "• Machine Learning và Deep Learning mang lại bước đột phá trong khám phá thuốc.")

# Slide 3: Mục tiêu nghiên cứu
add_slide("2. Mục tiêu nghiên cứu",
          "• Xây dựng mô hình lai CNN-BiLSTM (DDI-Hybrid).\n"
          "• CNN: Trích xuất đặc trưng không gian từ chuỗi SMILES.\n"
          "• BiLSTM: Nắm bắt ngữ cảnh dài và thông tin liên kết phức tạp.\n"
          "• Dự đoán chính xác 86 loại tương tác thuốc khác nhau.")

# Slide 4: Đối tượng và Phạm vi
add_slide("3. Đối tượng và Phạm vi nghiên cứu",
          "• Đối tượng: Bộ dữ liệu Gold Standard DDI Dataset từ DrugBank.\n"
          "• Phạm vi: Xử lý chuỗi SMILES của thuốc.\n"
          "• Mục tiêu: Phân loại hiện tượng tương tác khi phối hợp hai loại thuốc.")

# Slide 5: Kiến trúc mô hình
add_slide("4. Kiến trúc mô hình DDI-Hybrid",
          "• Feature Encoding: Mã hóa Morgan Fingerprints (2048 bit).\n"
          "• CNN Block: 4 lớp Conv1D (16, 32, 48, 64 bộ lọc).\n"
          "• BiLSTM Block: 2 lớp (128 và 96 đơn vị) xử lý hai chiều.\n"
          "• Output Layer: Lớp Dense (256) và Softmax phân loại 86 lớp.")

# Slide 6: Kết quả thực nghiệm
add_slide("5. Kết quả và So sánh",
          "• Accuracy đạt 95.38%.\n"
          "• Macro-F1 đạt 89.80% (vượt trội so với các mô hình Ensemble và Deep-DDI).\n"
          "• Khả năng nhận diện tốt các loại tương tác hiếm (lớp ít mẫu).")

# Slide 7: Tài nguyên và Kết luận
add_slide("6. Tài nguyên và Kết luận",
          "• Ngôn ngữ: Python 3.7.10 | Thư viện: TensorFlow, Keras, RDKit.\n"
          "• Mã nguồn: https://github.com/nguyenquypro17/project_ml.git\n"
          "• Mô hình có khả năng tổng quát hóa tốt, hỗ trợ bác sĩ kê đơn an toàn.")

# Lưu file
prs.save('BTL_HocMay_Nhom19.pptx')
print("Đã tạo file BTL_HocMay_Nhom19.pptx thành công!")